import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Thiết lập lề trang chuẩn (2cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Palette màu
    COLOR_PRIMARY = RGBColor(27, 94, 32)      # Xanh lá đậm #1B5E20
    COLOR_SECONDARY = RGBColor(46, 125, 50)  # Xanh lá vừa #2E7D32
    COLOR_ACCENT = RGBColor(230, 81, 0)      # Cam đậm #E65100
    COLOR_TEXT = RGBColor(33, 33, 33)        # Xám đen chữ #212121
    COLOR_MUTED = RGBColor(97, 97, 97)       # Xám phụ #616161

    # Style Tiêu Đề Chính (Title)
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("TÀI LIỆU ĐẶC TẢ TỔNG QUAN & CHI TIẾT TÍNH NĂNG\nGAME HOA QUẢ NỔI GIẬN")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(20)
    run_sub = sub_p.add_run("DỰ ÁN GAME PHÒNG THỦ CHIẾN THUẬT (PLANTS VS ZOMBIES CLONE) - C# WINFORMS .NET 8")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_ACCENT

    # Bảng thông tin dự án
    info_table = doc.add_table(rows=5, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Tên dự án:", "Hoa Quả Nổi Giận (Plants vs Zombies Remake)"),
        ("Nền tảng công nghệ:", "C# .NET 8 (Windows Forms), Guna.UI2 WinForms Framework"),
        ("Kiến trúc phần mềm:", "Mô hình phân tầng (Forms - Engine - Managers - Models)"),
        ("Đối tượng sử dụng:", "Báo cáo đồ án môn học, Tài liệu đặc tả chức năng (SRS) & Bàn giao"),
        ("Ngày lập tài liệu:", "Tháng 10 / 2026")
    ]
    for row_idx, (label, val) in enumerate(info_data):
        row = info_table.rows[row_idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.6)
        set_cell_background(cell_lbl, "E8F5E9")
        set_cell_background(cell_val, "FAFAFA")
        set_cell_margins(cell_lbl, 80, 80, 120, 120)
        set_cell_margins(cell_val, 80, 80, 120, 120)
        
        p1 = cell_lbl.paragraphs[0]
        r1 = p1.add_run(label)
        r1.font.name = "Arial"
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_PRIMARY
        
        p2 = cell_val.paragraphs[0]
        r2 = p2.add_run(val)
        r2.font.name = "Arial"
        r2.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_SECONDARY
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = COLOR_ACCENT
        return p

    def add_body(text, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Arial"
            r_pre.font.size = Pt(10)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_PRIMARY
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.italic = italic
        r.font.color.rgb = COLOR_TEXT
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Arial"
            r_pre.font.size = Pt(10)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_PRIMARY
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_TEXT
        return p

    # =========================================================================
    # PHẦN 1: GIỚI THIỆU TỔNG QUAN
    # =========================================================================
    add_heading_1("1. TỔNG QUAN DỰ ÁN GAME \"HOA QUẢ NỔI GIẬN\"")
    add_body("Game \"Hoa Quả Nổi Giận\" là một trò chơi chiến thuật phòng thủ thời gian thực (Tower Defense) được phát triển hoàn toàn trên nền tảng C# và Microsoft .NET 8 (Windows Forms), lấy cảm hứng từ tựa game huyền thoại kinh điển Plants vs. Zombies. Trò chơi đặt người chơi vào vai người làm vườn dũng cảm, sử dụng năng lượng ánh sáng mặt trời để trồng và bố trí các loài hoa quả chiến binh nhằm ngăn chặn các đợt xâm lăng của binh đoàn Zombie đông đảo.")
    
    add_heading_2("1.1. Mục tiêu và Điểm nổi bật")
    add_bullet(" Tái hiện chân thực cơ chế trò chơi huyền thoại với nhịp độ chơi nhanh, hấp dẫn, độ khó tăng dần qua 5 màn chơi.", "Gameplay Chuẩn Mực:")
    add_bullet(" Sử dụng bộ thư viện Guna.UI2 WinForms cao cấp kết hợp với kỹ thuật đồ họa vector GDI+ chống răng cưa (Anti-Aliasing) mịn màng, đổ bóng 3D, bo tròn góc hiện đại.", "Giao Diện Đồ Họa Đẳng Cấp:")
    add_bullet(" Hỗ trợ quản lý nhiều người chơi độc lập. Mỗi người chơi được lưu riêng tiến trình màn chơi, số sao đạt được, điểm tích lũy và thời gian chơi gần nhất.", "Hệ Thống Hồ Sơ (Profile) Độc Lập:")
    add_bullet(" Toàn bộ màn hình đều được xây dựng chuẩn WinForms Designer dạng điều khiển kéo thả (Drag-and-Drop), cho phép lập trình viên và người dùng dễ dàng mở bằng Visual Studio để chỉnh vị trí các nhân vật, nút bấm bằng chuột.", "Khả Năng Kéo Thả Trực Quan (Designer):")
    add_bullet(" Cơ chế nạp ảnh thông minh ImageLoader tự động nhận diện file ảnh bên ngoài (PNG), không khóa file (No File Locking), tự động đồng bộ khi thay đổi.", "Tùy Biến Tài Nguyên Linh Hoạt:")

    add_heading_2("1.2. Kiến trúc Module Dự Án")
    add_body("Dự án được phân chia rõ ràng theo các tầng kiến trúc chuẩn:")
    add_bullet(" Chứa toàn bộ các Form giao diện người dùng: frmSplash, frmMenu, frmGame, frmPause, frmResult, frmProgress, frmScoreboard, frmLevelSelect, frmHelp, frmSettings, frmConfirmExit, frmNewPlayer.", "1. Forms (Giao diện):")
    add_bullet(" Trái tim vận hành game, bao gồm GameEngine (Vòng lặp 60 FPS, xử lý va chạm, đạn, máu), GamePanel (Khung vẽ đồ họa sân vườn 5x9, thanh tiến trình, thanh chọn hạt giống), WaveManager (quản lý đợt quái) và GameConfig (hằng số hệ thống).", "2. Engine (Động cơ Game):")
    add_bullet(" Các lớp nghiệp vụ quản trị dữ liệu: ImageLoader (xử lý hình ảnh), SoundManager (âm thanh nhạc nền & SFX), ProfileManager (tiến trình người chơi), ScoreManager (bảng thành tích Top 10), SettingsManager (cấu hình).", "3. Managers (Quản lý dịch vụ):")
    add_bullet(" Định nghĩa các thực thể trong game: Các loài Hoa quả (Banana, Grape, Coconut, Ice Strawberry, Chili, Watermelon, Tomato, Durian), các loài Zombie (Normal, Cone, Bucket, Runner, Shield, Boss), Đạn (Bullet), Xe cắt cỏ (LawnMower), Mặt trời (SunToken)...", "4. Models (Thực thể & Dữ liệu):")

    # =========================================================================
    # PHẦN 2: DANH MỤC CÁC MÀN HÌNH VÀ GIAO DIỆN CHÍNH
    # =========================================================================
    add_heading_1("2. CHI TIẾT TỪNG MÀN HÌNH & GIAO DIỆN (FORMS)")
    add_body("Hệ thống bao gồm 12 Form giao diện hoàn chỉnh, được liên kết logic chặt chẽ với nhau:")

    add_heading_2("2.1. Màn hình Chờ Khởi động (frmSplash)")
    add_bullet(" Hiển thị khi bắt đầu mở game, tải trước tài nguyên hình ảnh và âm thanh vào bộ nhớ cache, tạo cảm giác chuyên nghiệp như các game thương mại.", "Chức năng:")
    add_bullet(" Thanh tiến trình Guna2ProgressBar chạy từ 0% đến 100% kèm nhãn trạng thái \"Đang tải tài nguyên...\".", "Thanh tải (Loading):")
    add_bullet(" Nửa trên hiển thị tiêu đề chính \"HOA QUẢ NỔI GIẬN\" màu vàng kim rực rỡ và khẩu hiệu sân vườn.", "Nửa trên:")
    add_bullet(" Phía bên trái là đội hình Hoa Quả (Nho, Chuối, Dừa, Dâu tuyết), chính giữa là huy hiệu đối đầu VS tròn lửa, phía bên phải là đội hình Zombie (Zombie nón, Zombie xô, Zombie Boss, Zombie thường).", "Nửa dưới (Chiến trường đối đầu):")
    add_bullet(" Toàn bộ các nhân vật trên frmSplash đều là các PictureBox kéo thả độc lập trong Visual Studio Designer, cho phép người dùng tùy ý di chuyển vị trí, chỉnh to nhỏ bằng chuột.", "Đặc điểm kéo thả Designer:")

    add_heading_2("2.2. Màn hình Menu Chính (frmMenu)")
    add_bullet(" Trung tâm điều hướng toàn bộ trò chơi, phát nhạc nền Menu rộn rã.", "Chức năng:")
    add_bullet(" Thanh nhãn thông tin người chơi hiện tại: \"👤 Người chơi: [Tên] | Màn chơi hiện tại: [Tên Màn] | Điểm: [Điểm] | ⭐: [Số Sao]\".", "Dải thông tin Người chơi:")
    add_bullet(" Mở hộp thoại nhập tên người chơi mới và bắt đầu chơi ngay từ Màn 1.", "Nút 'CHƠI MỚI':")
    add_bullet(" Mở màn hình danh sách người chơi để theo dõi tiến trình và chọn tiếp tục chơi phần của mình.", "Nút 'TIẾN TRÌNH':")
    add_bullet(" Cho phép chọn bất kỳ màn chơi nào đã mở khóa (từ Màn 1 đến Màn 5).", "Nút 'CHỌN MÀN CHƠI':")
    add_bullet(" Mở bảng vinh danh Top 10 kỷ lục điểm cao nhất.", "Nút 'BẢNG THÀNH TÍCH':")
    add_bullet(" Mở cẩm nang hướng dẫn cách chơi và thông tin các loài cây, quái vật.", "Nút 'HƯỚNG DẪN':")
    add_bullet(" Mở hộp thoại cài đặt bật/tắt nhạc nền và hiệu ứng âm thanh.", "Nút 'CÀI ĐẶT ÂM THANH':")
    add_bullet(" Mở hộp thoại xác nhận thoát game an toàn.", "Nút 'THOÁT GAME':")

    add_heading_2("2.3. Hộp thoại Nhập Người Chơi Mới (frmNewPlayer)")
    add_bullet(" Cho phép nhập tên người chơi mới khi bấm \"CHƠI MỚI\" hoặc bấm thêm người chơi trong bảng tiến trình.", "Chức năng:")
    add_bullet(" TextBox bo góc Guna2, tự động Focus con trỏ chuột, hỗ trợ gõ phím Enter để xác nhận nhanh.", "Điều khiển:")
    add_bullet(" Tự động gán tên mặc định nếu người chơi để trống, khởi tạo hồ sơ lưu trữ với Màn 1 mở khóa, điểm = 0, sao = 0.", "Cơ chế xử lý:")

    add_heading_2("2.4. Màn hình Quản Lý Tiến Trình (frmProgress)")
    add_bullet(" Quản lý tập trung mọi hồ sơ người chơi đa người dùng (Multi-Profile).", "Chức năng:")
    add_bullet(" Bảng DataGridView thiết kế hiện đại hiển thị: STT, Tên Người Chơi, Màn Chơi Hiện Tại, Màn Mở Khóa, Điểm Tích Lũy, Số Sao Đạt Được, Thời Gian Chơi Gần Nhất.", "Danh sách Hồ sơ:")
    add_bullet(" Bấm để vào chơi ngay màn chơi đang dở của người được chọn (hoặc có thể nhấp đúp trực tiếp vào dòng người chơi).", "Nút 'CHƠI TIẾP':")
    add_bullet(" Chuyển sang màn hình chọn màn với người chơi đang chọn.", "Nút 'CHỌN MÀN':")
    add_bullet(" Xóa hồ sơ người chơi được chọn (kèm hộp thoại xác nhận cảnh báo an toàn).", "Nút 'XÓA TIẾN TRÌNH':")
    add_bullet(" Mở hộp thoại tạo thêm nhân vật chơi mới.", "Nút 'TẠO MỚI':")
    add_bullet(" Trở lại màn hình Menu chính.", "Nút 'QUAY LẠI':")

    add_heading_2("2.5. Màn hình Chọn Màn Chơi (frmLevelSelect)")
    add_bullet(" Trực quan hóa bản đồ chiến dịch 5 màn chơi.", "Chức năng:")
    add_bullet(" 5 nút đại diện cho 5 màn chơi. Các màn đã mở khóa sẽ có màu vàng cam tươi sáng kèm số sao đạt được (⭐⭐⭐); các màn chưa chơi tới sẽ tự động bị khóa (màu xám, có biểu tượng 🔒 ĐÃ KHÓA và không thể bấm).", "Trạng thái Màn:")
    add_bullet(" Màn 1 (Vườn Sau Nhà), Màn 2 (Vườn Trước Nhà), Màn 3 (Vườn Cây Ăn Trái), Màn 4 (Khu Vườn Đêm), Màn 5 (Trận Cuối Cùng).", "Các màn chơi:")

    add_heading_2("2.6. Màn hình Trận Đấu Chính (frmGame & GamePanel)")
    add_bullet(" Khung cửa sổ kích thước chuẩn HD chứa canvas trò chơi thời gian thực (60 FPS).", "Chức năng:")
    add_bullet(" Lưới sân cỏ 5 hàng x 9 cột chuẩn quốc tế. Mỗi ô có kích thước 96 x 100 pixel.", "Khung sân vườn:")
    add_bullet(" Nằm ở mép trên cùng, gồm ô đếm tổng số Nắng hiện có, danh sách thẻ bài hạt giống hoa quả kèm giá nắng và hiệu ứng mờ dần hồi chiêu (Cooldown Overlay), biểu tượng Xẻng đào cây.", "Thanh Chọn Hạt Giống (Seed Chooser):")
    add_bullet(" Nằm ở cạnh dưới cùng sân, hiển thị tiến độ đợt zombie tiến tới. Có cắm cờ đỏ 🚩 tại vị trí đợt tấn công lớn (Huge Wave / Final Wave). Đã loại bỏ biểu tượng cảnh báo ⚠️ theo yêu cầu để giao diện thanh thoát.", "Thanh Tiến Trình Đợt Quái (Bottom Wave Bar):")
    add_bullet(" Đầu mỗi hàng cỏ đều có 1 chiếc máy cắt cỏ màu đỏ bảo vệ. Tự động lao thẳng tiêu diệt toàn bộ zombie trên hàng nếu có zombie vượt qua vạch phòng ngự cuối cùng.", "Xe cắt cỏ (Lawn Mower):")
    add_bullet(" Tên màn chơi và nút Tạm Dừng [⏸ TẠM DỪNG] được đặt gọn gàng ở góc trên, không đè lấn lên các ô trồng cây hay hoa quả hàng cuối cùng.", "Bố cục đỉnh màn:")

    add_heading_2("2.7. Hộp thoại Tạm Dừng (frmPause)")
    add_bullet(" Mở ra khi người chơi bấm nút Tạm Dừng hoặc bấm phím ESC trong trận đấu. Đóng băng Game Loop.", "Chức năng:")
    add_bullet(" [TIẾP TỤC] quay lại chơi tiếp; [CHƠI LẠI] tải lại màn đấu từ đầu; [CÀI ĐẶT] chỉnh âm lượng; [MENU CHÍNH] thoát về trang chủ.", "Các nút điều khiển:")

    add_heading_2("2.8. Màn hình Kết Quả Trận Đấu (frmResult)")
    add_bullet(" Xuất hiện khi kết thúc màn chơi (Thắng hoặc Thua).", "Chức năng:")
    add_bullet(" Tiêu đề lớn rực rỡ \"CHIẾN THẮNG!\" hoặc \"THẤT BẠI!\", biểu tượng số sao ⭐⭐⭐ đạt được (tính theo số xe cắt cỏ còn nguyên vẹn và máu cây cối), tổng điểm thưởng kiếm được.", "Thông tin hiển thị:")
    add_bullet(" Nút [CHƠI TIẾP] (mở màn tiếp theo nếu thắng), nút [CHƠI LẠI], và nút [QUAY VỀ MENU].", "Lựa chọn hành động:")

    add_heading_2("2.9. Bảng Thành Tích Xếp Hạng (frmScoreboard)")
    add_bullet(" Vinh danh kỷ lục điểm cao nhất (Top 10 High Scores).", "Chức năng:")
    add_bullet(" Tự động tô màu nổi bật huy chương: Top 1 (Vàng óng), Top 2 (Bạc sáng), Top 3 (Đồng ánh kim).", "Vinh danh huy chương:")
    add_bullet(" Cột hiển thị: Thứ hạng, Tên người chơi, Màn đạt được, Tổng điểm, Số sao, Thời gian sống sót, Ngày lưu.", "Cột dữ liệu:")
    add_bullet(" Nút [XÓA THÀNH TÍCH] (có hỏi xác nhận) và nút [QUAY LẠI].", "Các nút:")

    add_heading_2("2.10. Cẩm Nang Hướng Dẫn & Bách Khoa Toàn Thư (frmHelp)")
    add_bullet(" Cung cấp toàn bộ tri thức trong game thông qua 4 Tab Guna2 chuyển đổi mượt mà:", "Chức năng:")
    add_bullet(" Hướng dẫn chi tiết cách nhặt năng lượng nắng, trồng hoa quả, phòng thủ, sử dụng xẻng và cơ chế xe cắt cỏ.", "Tab 1 (CÁCH CHƠI):")
    add_bullet(" Thông tin chi tiết hình ảnh, chỉ số máu, sát thương, giá nắng, công dụng của toàn bộ 8 loài cây.", "Tab 2 (HOA QUẢ):")
    add_bullet(" Thông tin chi tiết hình ảnh, máu giáp, tốc độ, mức độ nguy hiểm của toàn bộ 6 loài Zombie.", "Tab 3 (ZOMBIE):")
    add_bullet(" Thông tin phiên bản, đồ án môn học, công nghệ phát triển.", "Tab 4 (THÔNG TIN):")

    add_heading_2("2.11. Cài Đặt Âm Thanh (frmSettings)")
    add_bullet(" Quản lý âm thanh tập trung, gọn gàng và thẩm mỹ.", "Chức năng:")
    add_bullet(" Công tắc gạt (Toggle Switch) bật/tắt Hiệu ứng âm thanh (Sound Effects) và Nhạc nền (Background Music). Đã lược bỏ thanh trượt âm lượng tổng theo yêu cầu để giao diện tối giản.", "Các tùy chọn:")

    add_heading_2("2.12. Hộp thoại Xác Nhận Thoát Game (frmConfirmExit)")
    add_bullet(" Thiết kế pop-up cảnh báo hiện đại với thông điệp: \"Bạn có chắc chắn muốn rời khỏi vườn không?\". Nút [THOÁT GAME] đỏ và nút [Ở LẠI CHƠI] xanh giúp chống bấm nhầm.", "Chức năng:")

    # =========================================================================
    # PHẦN 3: CƠ CHẾ GAMEPLAY & QUY TẮC TRÒ CHƠI
    # =========================================================================
    add_heading_1("3. CƠ CHẾ GAMEPLAY & QUY TẮC TRÒ CHƠI")
    
    add_heading_2("3.1. Hệ thống Năng Lượng Ánh Sáng (Nắng)")
    add_bullet(" Nắng là đơn vị tiền tệ duy nhất dùng để mua và trồng các loại hoa quả.", "Vai trò:")
    add_bullet(" Ban đầu người chơi được cấp sẵn 150 Nắng. Giới hạn tích lũy tối đa lên tới 9.990 Nắng.", "Khởi đầu & Giới hạn:")
    add_bullet(" Cứ mỗi 9 giây, một bông mặt trời sẽ rơi ngẫu nhiên từ bầu trời xuống sân cỏ.", "Rơi từ bầu trời:")
    add_bullet(" Cây Chuối Nắng sản sinh tia nắng định kỳ mỗi 8 - 12 giây.", "Cây Chuối sản sinh:")
    add_bullet(" Người chơi chỉ cần nhấp chuột trái vào bông mặt trời để thu hoạch. Mỗi lần nhặt được cộng ngay +50 Nắng (đã điều chỉnh tăng gấp đôi so với 25 trước đây). Nếu không nhặt, sau 8 giây bông nắng sẽ tự biến mất.", "Giá trị thu hoạch:")

    add_heading_2("3.2. Quy tắc Trồng Cây & Thời Gian Hồi Chiêu (Cooldown)")
    add_bullet(" Bấm chuột chọn thẻ hạt giống trên khay -> di chuột vào ô đất trống trên sân cỏ -> nhấp chuột trái để gieo trồng.", "Cách trồng:")
    add_bullet(" Mỗi ô đất trên lưới 5x9 chỉ được trồng tối đa 1 cây. Không thể trồng đè lên ô đã có cây trừ khi dùng Xẻng đào cây cũ đi.", "Quy tắc ô lưới:")
    add_bullet(" Sau khi trồng một cây, thẻ bài của cây đó sẽ chuyển sang trạng thái hồi chiêu với hiệu ứng mờ dần (Cooldown Overlay). Người chơi phải chờ hết thời gian hồi chiêu và tích đủ số Nắng yêu cầu mới có thể trồng cây tiếp theo cùng loại.", "Thời gian hồi chiêu:")

    add_heading_2("3.3. Cơ chế Tấn công & Tương tác Va chạm (Collision)")
    add_bullet(" Các loài hoa quả bắn đạn (Nho, Dâu tây tuyết, Dưa hấu) chỉ khai hỏa khi phát hiện có Zombie xuất hiện trên cùng hàng phía trước mặt nó. Khi không có zombie, cây sẽ nghỉ ngơi để tiết kiệm đạn.", "Cảm biến hàng:")
    add_bullet(" Đạn bay thẳng từ trái sang phải với tốc độ 450 px/giây. Khi chạm vào Zombie đầu tiên trên đường bay, đạn sẽ phát nổ, gây sát thương và biến mất.", "Đường bay của đạn:")
    add_bullet(" Zombie di chuyển chậm từ phải sang trái. Khi gặp hoa quả, Zombie sẽ dừng lại và liên tục cắn xé. Khi hoa quả hết máu và biến mất, Zombie sẽ tiếp tục tiến bước.", "Zombie cắn hoa quả:")
    add_bullet(" Thanh máu (HP Bar) của Zombie được thiết kế thông minh: Thanh máu nền đỏ sẽ ngắn dần chính xác theo phần trăm máu còn lại. Khi Zombie bị mất giáp (vỡ nón, vỡ xô, rơi khiên), máu sẽ tụt nhanh và thanh máu hiển thị chân thực.", "Hiển thị máu Zombie:")

    add_heading_2("3.4. Xe Cắt Cỏ (Lawn Mower) - Phòng tuyến Cuối cùng")
    add_bullet(" Có 5 chiếc xe cắt cỏ đặt ở vị trí cột đầu tiên bên trái của cả 5 hàng.", "Bố trí:")
    add_bullet(" Khi Zombie bất kỳ vượt qua mọi hoa quả phòng ngự và chạm vào vạch cuối cùng, xe cắt cỏ tương ứng ở hàng đó sẽ khởi động, phóng với tốc độ 400 px/giây dọn sạch toàn bộ zombie trên hàng đó.", "Kích hoạt:")
    add_bullet(" Mỗi xe cắt cỏ chỉ dùng được 1 lần duy nhất trong màn chơi. Sau khi đã phóng đi, hàng đó sẽ không còn lá chắn bảo vệ cuối cùng. Nếu Zombie vượt qua lần nữa, người chơi sẽ bị THẤT BẠI (Brain Eaten).", "Giới hạn:")

    # =========================================================================
    # PHẦN 4: HỆ THỐNG HOA QUẢ (PLANTS)
    # =========================================================================
    add_heading_1("4. DANH MỤC CHI TIẾT CÁC LOÀI HOA QUẢ (PLANTS)")
    add_body("Bảng tổng hợp thông số và kỹ năng của toàn bộ 8 loài hoa quả chiến binh:")

    plant_table = doc.add_table(rows=9, cols=6)
    plant_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    plant_headers = ["Tên Hoa Quả", "Giá Nắng", "Máu (HP)", "Sát Thương", "Hồi Chiêu", "Kỹ Năng & Vai Trò Đặc Biệt"]
    for i, h in enumerate(plant_headers):
        cell = plant_table.rows[0].cells[i]
        set_cell_background(cell, "1B5E20")
        set_cell_margins(cell, 100, 100, 80, 80)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    plant_data = [
        ("Chuối Nắng\n(BananaSun)", "50", "300", "0", "5.0s", "Sản sinh +50 Nắng định kỳ mỗi 8 - 12s. Nguồn kinh tế cốt lõi."),
        ("Nho Bắn Súng\n(GrapeShooter)", "100", "300", "20 / viên", "6.0s", "Bắn hạt nho liên tục (1.4s/viên) thẳng hàng. Xạ thủ cơ bản."),
        ("Dâu Tuyết\n(IceStrawberry)", "175", "300", "20 / viên", "7.5s", "Bắn đạn băng làm chậm 50% tốc độ di chuyển và tốc độ cắn của Zombie."),
        ("Dừa Hộ Vệ\n(CoconutWall)", "50", "4000", "0", "20.0s", "Tấm khiên siêu bền, máu gấp hơn 13 lần cây thường. Chặn đứng bước tiến của quái."),
        ("Ớt Lửa Hủy Diệt\n(ChiliBomb)", "125", "300", "1800", "25.0s", "Phát nổ thiêu rụi toàn bộ Zombie trên cùng 1 hàng chỉ trong 0.8s."),
        ("Dưa Hấu Đại Bác\n(Watermelon)", "300", "300", "80 + AoE", "15.0s", "Bắn quả dưa cầu vồng bay bổng, phát nổ gây 80 ST mục tiêu và 26 ST lan xung quanh."),
        ("Bẫy Cà Chua\n(TomatoMine)", "25", "300", "1800", "18.0s", "Chôn mình dưới đất mất 12s kích hoạt, Zombie dẫm phải sẽ phát nổ tức thì."),
        ("Sầu Riêng Gai\n(DurianThorn)", "100", "500", "20 / 0.5s", "8.0s", "Đặt dưới đất, Zombie dẫm qua bị đâm liên tục. Không thể bị Zombie cắn ăn.")
    ]

    for row_idx, row_data in enumerate(plant_data):
        row = plant_table.rows[row_idx + 1]
        bg_col = "F1F8E9" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, 80, 80, 80, 80)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # =========================================================================
    # PHẦN 5: HỆ THỐNG ZOMBIE
    # =========================================================================
    add_heading_1("5. DANH MỤC CHI TIẾT CÁC LOÀI ZOMBIE")
    add_body("Bảng thông số máu, tốc độ và mức độ nguy hiểm của toàn bộ 6 chủng loài Zombie (đã cân chỉnh máu hợp lý để bắn nhanh chết, trải nghiệm đã tay):")

    zombie_table = doc.add_table(rows=7, cols=6)
    zombie_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    zombie_headers = ["Chủng Zombie", "Máu Giáp", "Máu Thân", "Tốc Độ", "Điểm Hạ", "Đặc Điểm & Cách Khắc Chế"]
    for i, h in enumerate(zombie_headers):
        cell = zombie_table.rows[0].cells[i]
        set_cell_background(cell, "B71C1C")
        set_cell_margins(cell, 100, 100, 80, 80)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    zombie_data = [
        ("Zombie Thường\n(Normal Zombie)", "0", "80 HP", "28 px/s", "10", "Chậm chạp, mặc áo vest rách. Bị tiêu diệt sau khoảng 4 phát đạn Nho."),
        ("Zombie Nón Lá\n(Conehead Zombie)", "120 HP", "80 HP", "28 px/s", "20", "Đội nón bảo hộ giao thông. Phải bắn vỡ nón trước rồi mới bắn vào thân."),
        ("Zombie Xô Sắt\n(Buckethead)", "240 HP", "80 HP", "28 px/s", "30", "Đội xô thiếc kim loại siêu trâu. Nên dùng Dây tây tuyết làm chậm hoặc Ớt lửa."),
        ("Zombie Điền Kinh\n(Runner Zombie)", "0", "70 HP", "55 px/s", "25", "Chạy rất nhanh (gấp đôi bình thường), lao tới áp sát. Cần Dừa chặn hoặc Bẫy cà chua."),
        ("Zombie Khiên Gỗ\n(Shield Zombie)", "140 HP", "80 HP", "24 px/s", "40", "Cầm cánh cửa gỗ che chắn toàn bộ đạn thẳng. Bị khắc chế mạnh bởi Dưa hấu nổ lan."),
        ("Zombie Boss\n(Gargantuar Boss)", "0", "1800 HP", "20 px/s", "500", "Khổng lồ, mắt đỏ rực. Đập bẹp hoa quả, triệu hồi Zombie con. Cần dồn toàn lực Ớt & Dưa hấu.")
    ]

    for row_idx, row_data in enumerate(zombie_data):
        row = zombie_table.rows[row_idx + 1]
        bg_col = "FFEBEE" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, 80, 80, 80, 80)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(183, 28, 28)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # =========================================================================
    # PHẦN 6: CHIẾN DỊCH 5 MÀN CHƠI
    # =========================================================================
    add_heading_1("6. HỆ THỐNG CHIẾN DỊCH 5 MÀN CHƠI (LEVELS)")
    add_body("Mỗi màn chơi được thiết kế bối cảnh, tỷ lệ xuất hiện của các đợt Zombie và điều kiện hoàn thành riêng biệt:")

    add_heading_2("Màn 1: Vườn Sau Nhà (Day Garden Easy)")
    add_bullet(" Bối cảnh sân cỏ xanh ban ngày, nhịp độ êm đềm để làm quen.", "Bối cảnh:")
    add_bullet(" 2 Đợt sóng Zombie (Waves). Chỉ xuất hiện Zombie Thường và Zombie Nón Lá số lượng ít.", "Kẻ địch:")
    add_bullet(" Trồng 1-2 Chuối Nắng, xây 1 hàng Nho Bắn Súng là dễ dàng vượt qua.", "Mục tiêu:")

    add_heading_2("Màn 2: Vườn Trước Nhà (Front Yard Normal)")
    add_bullet(" Sân vườn thoáng đãng, tốc độ quái bắt đầu xuất hiện dồn dập hơn.", "Bối cảnh:")
    add_bullet(" 3 Đợt sóng. Xuất hiện thêm Zombie Điền Kinh (Runner) chạy nhanh bất ngờ.", "Kẻ địch:")
    add_bullet(" Cần mở khóa Dừa Hộ Vệ và Dâu Tuyết để chặn đứng các con chạy nhanh.", "Chiến thuật:")

    add_heading_2("Màn 3: Vườn Cây Ăn Trái (Orchard Defense)")
    add_bullet(" Sân cỏ rậm rạp cây trái, đợt bão Zombie tấn công quy mô lớn.", "Bối cảnh:")
    add_bullet(" 4 Đợt sóng. Xuất hiện Zombie Xô Sắt (Buckethead) và Zombie Khiên Gỗ (Shield).", "Kẻ địch:")
    add_bullet(" Bắt buộc phải phối hợp Dâu Tuyết làm chậm và Bẫy Cà Chua tiêu diệt quái trâu.", "Chiến thuật:")

    add_heading_2("Màn 4: Khu Vườn Đêm (Night Garden Hard)")
    add_bullet(" Màn đêm buông xuống, không còn nắng rơi từ bầu trời (chỉ trông cậy vào Chuối).", "Bối cảnh:")
    add_bullet(" 5 Đợt sóng liên tục, tốc độ xuất hiện cực nhanh với sự kết hợp của mọi chủng loại quái.", "Kẻ địch:")
    add_bullet(" Tận dụng Ớt Lửa (ChiliBomb) và Dưa Hấu Đại Bác để giải phóng hàng bị vây hãm.", "Chiến thuật:")

    add_heading_2("Màn 5: Trận Cuối Cùng (Final Boss Battle)")
    add_bullet(" Chiến trường sinh tử đỉnh cao, sấm chớp rền vang.", "Bối cảnh:")
    add_bullet(" 6 Đợt sóng khốc liệt. Ở đợt Final Wave, Zombie Boss Khổng Lồ xuất hiện với 1800 Máu.", "Kẻ địch:")
    add_bullet(" Dồn toàn bộ hỏa lực Dưa Hấu, thả Ớt Lửa chính xác và duy trì Dừa phòng thủ để tiêu diệt Boss.", "Chiến thuật:")

    add_heading_2("Cơ chế Đánh Giá Xếp Hạng 1 - 3 Sao:")
    add_bullet("⭐⭐⭐ (3 Sao Hoàn Hảo): Vượt qua màn chơi mà không để mất bất kỳ xe cắt cỏ nào và giữ được trên 70% số hoa quả.", "Xếp hạng 3 Sao:")
    add_bullet("⭐⭐ (2 Sao Xuất Sắc): Còn nguyên ít nhất 3 chiếc xe cắt cỏ khi kết thúc màn.", "Xếp hạng 2 Sao:")
    add_bullet("⭐ (1 Sao Đạt Chuẩn): Vượt qua màn chơi thành công dù đã kích hoạt hết xe cắt cỏ.", "Xếp hạng 1 Sao:")

    # =========================================================================
    # PHẦN 7: QUẢN LÝ DỮ LIỆU & LƯU TIẾN TRÌNH
    # =========================================================================
    add_heading_1("7. HỆ THỐNG LƯU TRỮ DỮ LIỆU & HỒ SƠ NGƯỜI CHƠI")
    add_body("Toàn bộ dữ liệu được quản lý tự động trong thư mục Data/, người chơi không bao giờ bị mất thành tích khi tắt máy hoặc thoát game giữa chừng:")

    add_bullet(" Lưu danh sách toàn bộ hồ sơ người chơi dưới định dạng JSON có cấu trúc. Lưu trữ: Id, Tên người chơi, Màn đang chơi (CurrentLevel), Màn cao nhất đã mở khóa (UnlockedLevel), Tổng điểm tích lũy (TotalScore), Danh sách số sao từng màn, Thời gian chơi cuối cùng (LastPlayed).", "1. profiles.json (Tiến trình người chơi):")
    add_bullet(" Lưu bảng vàng Top 10 kỷ lục điểm cao nhất dưới định dạng CSV phân tách dấu phẩy. Mỗi bản ghi gồm: Thứ hạng, Tên, Màn đạt được, Điểm số, Số sao, Thời gian sống sót (giây), Ngày giờ đạt kỷ lục.", "2. scores.csv (Bảng thành tích):")
    add_bullet(" Lưu cấu hình tùy chọn âm thanh của người chơi: SoundEnabled (Bật/Tắt SFX), MusicEnabled (Bật/Tắt Nhạc nền BGM).", "3. settings.ini (Cài đặt):")
    add_bullet(" Tự động ghi lại các trường hợp ngoại lệ (Exception) nếu có sự cố xảy ra, giúp nhà phát triển dễ dàng truy vết và sửa lỗi.", "4. error.log (Nhật ký lỗi hệ thống):")

    # =========================================================================
    # PHẦN 8: HỆ THỐNG ĐỒ HỌA & ÂM THANH
    # =========================================================================
    add_heading_1("8. HỆ THỐNG ĐỒ HỌA & ÂM THANH (GRAPHICS & AUDIO)")

    add_heading_2("8.1. Đồ Họa & Nạp Ảnh Thông Minh (ImageLoader)")
    add_bullet(" Hệ thống tự động sinh ra bộ ảnh vector GDI+ độ nét cao với phong cách Plants vs Zombies đáng yêu, mắt to long lanh, viền nổi 3D.", "Bộ ảnh vector tích hợp:")
    add_bullet(" Cho phép người dùng tùy ý chép các file ảnh định dạng PNG bên ngoài vào thư mục Assets/Images/ (ví dụ: fruit_banana.png, zombie_cone.png...). Game sẽ tự động ưu tiên lấy ảnh ngoài này thay cho ảnh mặc định.", "Hỗ trợ ảnh ngoài (PNG):")
    add_bullet(" Nạp ảnh qua bộ nhớ đệm MemoryStream, không dùng Image.FromFile giúp file ảnh trên đĩa cứng không bao giờ bị Windows khóa (No File Locking). Bạn có thể đổi hình, dán đè file thoải mái ngay khi game đang bật.", "Cơ chế không khóa file:")
    add_bullet(" Toàn bộ các cửa sổ Form đều được nạp file app_icon.ico chuẩn Windows, hiển thị biểu tượng Hoa Mặt Trời vàng rực trên thanh tiêu đề và thanh Taskbar.", "Icon cửa sổ chuẩn:")

    add_heading_2("8.2. Hệ Thống Âm Thanh (SoundManager)")
    add_bullet(" Nhạc nền menu (bgm_menu.wav) rộn ràng, nhạc nền chiến đấu sân vườn (bgm_game.wav) kịch tính, lôi cuốn.", "Nhạc nền (BGM):")
    add_bullet(" Tiếng bấm nút (sfx_click), tiếng bắn hạt (sfx_shoot), tiếng zombie cắn ngoạm (sfx_chomp), tiếng nổ thiêu rụi (sfx_explosion), tiếng nhặt hoa mặt trời (sfx_sun), tiếng đạn băng trúng đích (sfx_freeze), tiếng còi chiến thắng (sfx_win), tiếng chuông thất bại (sfx_lose).", "Hiệu ứng âm thanh (SFX):")

    # =========================================================================
    # PHẦN 9: HƯỚNG DẪN KÉO THẢ DESIGNER & TÙY BIẾN
    # =========================================================================
    add_heading_1("9. HƯỚNG DẪN KÉO THẢ DESIGNER & TÙY BIẾN TRONG VISUAL STUDIO")
    add_body("Hệ thống đã được thiết kế sẵn sàng để người dùng chỉnh sửa trực quan mà không cần can thiệp sâu vào code:")

    add_heading_2("9.1. Kéo Thả Trực Quan Trên Visual Studio Designer")
    add_bullet(" Trong Visual Studio, mở cây thư mục Solution Explorer bên phải, nhấp đúp chuột vào bất kỳ Form nào (ví dụ frmSplash.cs, frmMenu.cs...).", "Bước 1: Mở Design:")
    add_bullet(" Nhấp chuột trái vào bất kỳ con Hoa Quả, Zombie, Nút bấm hoặc Chữ tiêu đề nào.", "Bước 2: Chọn phần tử:")
    add_bullet(" Giữ chuột và kéo di chuyển đến vị trí mới; hoặc kéo các điểm neo vuông ở 4 góc để phóng to, thu nhỏ kích thước trực quan.", "Bước 3: Kéo thả & Thay đổi kích thước:")
    add_bullet(" Bấm phím F4 để mở cửa sổ Properties, bạn có thể chỉnh chính xác tọa độ Location (X, Y) và kích thước Size (Width, Height).", "Bước 4: Chỉnh tọa độ chính xác:")

    add_heading_2("9.2. Tùy Chỉnh Thông Số Game Trong File GameConfig.cs")
    add_body("Để tăng giảm độ khó của game theo ý muốn, bạn chỉ cần mở file Engine/GameConfig.cs và sửa các con số:")
    add_bullet(" Sửa HP_ZOMBIE_NORMAL = 80, HP_ZOMBIE_CONE_HAT = 120, HP_ZOMBIE_BUCKET_HAT = 240... (Số càng nhỏ quái càng nhanh chết, số càng lớn quái càng trâu).", "Chỉnh máu Zombie:")
    add_bullet(" Sửa SUN_VALUE = 50 (Mỗi lần nhặt hoa mặt trời cộng bao nhiêu Nắng).", "Chỉnh lượng Nắng nhặt được:")
    add_bullet(" Sửa SKY_SUN_INTERVAL = 9.0f (Khoảng cách thời gian rơi nắng từ trời xuống sân cỏ).", "Chỉnh tần suất rơi Nắng:")

    # Lưu tài liệu
    output_path = r"G:\hoa qua\HoaQuaNoiGian\Tai_Lieu_Chi_Tiet_Chuc_Nang_Game_Hoa_Qua_Noi_Gian.docx"
    doc.save(output_path)
    print(f"SUCCESS: Document saved to {output_path}")

if __name__ == "__main__":
    create_document()
